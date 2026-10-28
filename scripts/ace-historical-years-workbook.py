"""Build the private ACE pickup-year review workbook from a reconcile run directory.

Usage: python3 scripts/ace-historical-years-workbook.py (needs openpyxl) <run_dir> <out.xlsx>
Reads review_sheet.csv, queue_*.csv, unmatched_evidence.csv, summary.json. Private output.
"""
import csv, json, os, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

run, out = sys.argv[1], sys.argv[2]

# The workbook holds real names: refuse a path git would track.
probe = os.path.dirname(os.path.abspath(out)) or "."
inside = subprocess.run(["git", "-C", probe, "rev-parse", "--is-inside-work-tree"], capture_output=True).returncode == 0
if inside and subprocess.run(["git", "-C", probe, "check-ignore", "-q", "--no-index", os.path.abspath(out)]).returncode != 0:
    sys.exit(f"{out} would not be gitignored. Write it under scripts/reports/ or outside the repo.")

def rows(name):
    with open(f"{run}/{name}", newline="", encoding="utf-8") as f:
        return list(csv.reader(f))

FILL = {
    "CONFIRMED_YEAR": "C6EFCE", "STRONG_CANDIDATE": "DDEBF7", "SUSPECTED_YEAR": "FFF2CC",
    "CONFLICTING_EVIDENCE": "F8CBAD", "UNKNOWN": "EDEDED",
}
WIDTH = {"member_name": 24, "current_big_name": 24, "fam_name": 16, "review_reason": 60, "supporting_sources": 60,
         "conflicting_sources": 50, "source_sheet_and_cell": 40, "review_flags": 36, "reviewer_decision": 22}
HIDDEN = {"ace_node_id", "fam_id", "current_big_node_id", "canonical_member_id", "source_version"}

wb = Workbook()
readme = wb.active
readme.title = "Read me"
summary = json.load(open(f"{run}/summary.json"))
lines = [
    "ACE historical pickup-year review (#670) — PRIVATE. Do not share outside ACE Chair / Webmaster.",
    "",
    "The live website tree is the authority for every person's fam and Big. This sheet only proposes pickup YEARS.",
    "Never edit member_name, fam, Big or id columns: the plan rejects any row whose identity no longer matches the live site.",
    "",
    "How to review: edit ONLY the 'All nodes' tab. Queue tabs are protected, read-only copies for finding rows.",
    "  1. Read confidence_class, recommended year/term, supporting and conflicting sources, review_reason.",
    "  2. Fill reviewer_decision (free text), reviewer_confirmed_year (e.g. 2021–22 or 2021), reviewer_confirmed_term (FA21 / SP22, or blank if only the year is known).",
    "  3. Set approved_for_backfill to yes only when you are sure. Leave blank or 'no' otherwise.",
    "  4. Export the 'All nodes' tab as CSV and run: npm run ace:years -- plan --review <that.csv> ... (dry run).",
    "",
    "Classes: CONFIRMED_YEAR = official pairing names this Little under the live Big; STRONG_CANDIDATE = explicit added_term / looser official match;",
    "SUSPECTED_YEAR = Master workbook year row only (see MASTER_UNRELIABLE_IN_FAM); CONFLICTING_EVIDENCE = held; UNKNOWN = no reliable evidence.",
    "FA23 = 2023–24 but SP23 = 2022–23. Never infer a year from the Big, generation or graduation.",
    "",
    f"Live nodes: {summary['live_nodes']} in {summary['families']} fams. " + ", ".join(f"{k}: {v}" for k, v in summary["classes"].items()),
]
for i, line in enumerate(lines, 1):
    readme.cell(row=i, column=1, value=line)
readme.column_dimensions["A"].width = 150

def sheet(title, data, review=True, locked=False):
    ws = wb.create_sheet(title)
    if locked:
        ws.protection.sheet = True
    for r in data:
        ws.append(r)
    header = data[0] if data else []
    for c, name in enumerate(header, 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = WIDTH.get(name, 14)
        if name in HIDDEN:
            ws.column_dimensions[letter].hidden = True
        if name.startswith("reviewer_") or name == "approved_for_backfill":
            cell.fill = PatternFill("solid", fgColor="FFE699")
    ws.freeze_panes = "E2" if review else "A2"
    if data:
        ws.auto_filter.ref = ws.dimensions
    if review and "confidence_class" in header:
        ci = header.index("confidence_class") + 1
        for r in range(2, len(data) + 1):
            v = ws.cell(row=r, column=ci).value
            if v in FILL:
                ws.cell(row=r, column=ci).fill = PatternFill("solid", fgColor=FILL[v])
        if "approved_for_backfill" in header:
            col = get_column_letter(header.index("approved_for_backfill") + 1)
            dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{max(2, len(data))}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    return ws

sheet("All nodes", rows("review_sheet.csv"))
for q, title in [("conflicting_years", "Conflicting years"), ("ambiguous_identity", "Ambiguous identity"),
                 ("chronology", "Chronology"), ("roots", "Roots"), ("missing_evidence", "Missing evidence")]:
    sheet(title, rows(f"queue_{q}.csv"), locked=True)
sheet("Unmatched source rows", rows("unmatched_evidence.csv"), review=False, locked=True)
per = [["fam", "live_nodes", "with_canonical_member", "roots", "same_name_groups", "source_matched", "legacy_only",
        "master_agreed", "master_compared"] + list(summary["classes"].keys())]
for f in summary["per_fam"]:
    per.append([f["fam_name"], f["live_nodes"], f["with_canonical_member"], f["roots"], f["same_name_groups"], f["source_matched"],
                f["legacy_only"], f["master_agreed"], f["master_compared"]] + [f["classes"][k] for k in summary["classes"]])
sheet("Per-fam summary", per, review=False, locked=True)
wb.save(out)
print("saved", out)
